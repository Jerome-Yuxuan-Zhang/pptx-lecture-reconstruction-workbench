from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from pptx.enum.shapes import PP_PLACEHOLDER

from app.models.ir_models import BlockStyle, BoundingBox, ContentBlock, TextRun


@dataclass(slots=True)
class ParsedTextParagraph:
    text: str
    level: int
    alignment: str | None
    runs: list[TextRun]
    bullet: bool


def _alignment_name(value: Any) -> str | None:
    if value is None:
        return None
    return getattr(value, "name", str(value)).lower()


def _run_color(run: Any) -> str | None:
    color = getattr(run.font.color, "rgb", None)
    return str(color) if color else None


def parse_text_paragraphs(shape: Any) -> tuple[list[ParsedTextParagraph], BlockStyle]:
    paragraphs: list[ParsedTextParagraph] = []
    font_sizes: list[float] = []
    font_names: list[str] = []
    placeholder_title = False

    if getattr(shape, "is_placeholder", False):
        placeholder_type = shape.placeholder_format.type
        placeholder_title = placeholder_type in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}

    for paragraph in shape.text_frame.paragraphs:
        runs: list[TextRun] = []
        paragraph_text_parts: list[str] = []
        for run in paragraph.runs:
            text = run.text or ""
            paragraph_text_parts.append(text)
            font_size = run.font.size.pt if run.font.size else None
            font_name = run.font.name or None
            if font_size:
                font_sizes.append(font_size)
            if font_name:
                font_names.append(font_name)
            runs.append(
                TextRun(
                    text=text,
                    bold=bool(run.font.bold),
                    italic=bool(run.font.italic),
                    underline=bool(run.font.underline),
                    color=_run_color(run),
                    font_size_pt=font_size,
                    font_name=font_name,
                    level=paragraph.level,
                    superscript=bool(getattr(run.font, "superscript", False)),
                    subscript=bool(getattr(run.font, "subscript", False)),
                )
            )
        paragraph_text = "".join(paragraph_text_parts).strip()
        if not paragraph_text and paragraph.text:
            paragraph_text = paragraph.text.strip()
        if not paragraph_text:
            continue
        bullet = paragraph.level > 0 or paragraph_text[:1] in {"•", "-", "–", "·"}
        paragraphs.append(
            ParsedTextParagraph(
                text=paragraph_text,
                level=paragraph.level,
                alignment=_alignment_name(paragraph.alignment),
                runs=runs or [TextRun(text=paragraph_text, level=paragraph.level)],
                bullet=bullet,
            )
        )

    style = BlockStyle(
        font_size_pt=max(font_sizes) if font_sizes else None,
        font_name=font_names[0] if font_names else None,
        alignment=paragraphs[0].alignment if paragraphs else None,
        bullet=any(item.bullet for item in paragraphs),
        level=min((item.level for item in paragraphs), default=0),
        is_placeholder_title=placeholder_title,
        shape_type=getattr(shape.shape_type, "name", str(shape.shape_type)),
    )
    return paragraphs, style


def build_text_block(shape: Any, block_id: str, z_index: int) -> ContentBlock | None:
    paragraphs, style = parse_text_paragraphs(shape)
    if not paragraphs:
        return None

    bbox = BoundingBox(
        x=shape.left.pt, y=shape.top.pt, width=shape.width.pt, height=shape.height.pt
    )
    return ContentBlock(
        block_id=block_id,
        block_type="text",
        bbox=bbox,
        z_index=z_index,
        semantic_role="body",
        content={
            "text": "\n".join(paragraph.text for paragraph in paragraphs),
            "paragraphs": [
                {
                    "text": paragraph.text,
                    "level": paragraph.level,
                    "alignment": paragraph.alignment,
                    "runs": [run.model_dump() for run in paragraph.runs],
                    "bullet": paragraph.bullet,
                }
                for paragraph in paragraphs
            ],
        },
        style=style,
        source_shape_name=shape.name,
    )
